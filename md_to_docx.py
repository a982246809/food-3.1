import sys
import re

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "-i", "https://pypi.tuna.tsinghua.edu.cn/simple"])
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
    from docx.oxml.ns import qn

def md_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        md_text = f.read()

    doc = Document()
    # 设置全文字体为宋体
    doc.styles['Normal'].font.name = u'宋体'
    doc.styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), u'宋体')
    doc.styles['Normal'].font.size = Pt(12)

    lines = md_text.split('\n')
    
    in_code_block = False
    in_table = False
    table_data = []

    for line in lines:
        stripped = line.strip()
        
        if stripped.startswith('```'):
            in_code_block = not in_code_block
            continue
            
        if in_code_block:
            p = doc.add_paragraph()
            run = p.add_run(line) # Keep original indentation
            run.font.name = 'Consolas'
            run.font.size = Pt(10)
            continue
            
        if stripped.startswith('|'):
            in_table = True
            # Parse table row
            cells = [c.strip() for c in stripped.split('|') if c.strip() or c == '']
            if cells:
                # remove empty ends caused by split
                if stripped.startswith('|'): cells = cells[1:]
                if stripped.endswith('|'): cells = cells[:-1]
                
                # Check if it's separator row
                if all(c.replace('-', '').replace(':', '') == '' for c in cells):
                    continue
                table_data.append(cells)
            continue
        elif in_table and not stripped.startswith('|'):
            in_table = False
            # Draw table
            if table_data:
                rows = len(table_data)
                cols = max(len(r) for r in table_data)
                table = doc.add_table(rows=rows, cols=cols)
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < cols:
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_text
                doc.add_paragraph() # Add space after table
            table_data = []

        if not stripped:
            continue

        if stripped.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[4:])
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.name = u'黑体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), u'黑体')
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(stripped[3:])
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.name = u'黑体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), u'黑体')
        elif stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.add_run(stripped[2:])
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.name = u'黑体'
            run._element.rPr.rFonts.set(qn('w:eastAsia'), u'黑体')
        elif stripped.startswith('- '):
            p = doc.add_paragraph(stripped[2:], style='List Bullet')
        else:
            # Handle bold text within paragraph
            p = doc.add_paragraph()
            segments = re.split(r'(\*\*.*?\*\*)', stripped)
            for seg in segments:
                if seg.startswith('**') and seg.endswith('**'):
                    run = p.add_run(seg[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(seg)
            
    # If file ends with table
    if table_data:
        rows = len(table_data)
        cols = max(len(r) for r in table_data)
        table = doc.add_table(rows=rows, cols=cols)
        table.style = 'Table Grid'
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_text in enumerate(row):
                if c_idx < cols:
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_text

    doc.save(docx_path)

if __name__ == '__main__':
    md_to_docx(sys.argv[1], sys.argv[2])
